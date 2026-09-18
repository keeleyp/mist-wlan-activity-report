#!/usr/bin/env python3
"""
mist_wlan_activity_report.py — Juniper Mist WLAN Coverage + 5GHz Client Activity Report
==========================================================================================
Combines two things per AP:

  1. Which SSIDs actually reach it, resolved the same way Mist pushes WLAN Templates
     down to APs: a WLAN belongs to a Template (wlan.template_id); a Template applies
     to a set of sites (applies minus exceptions); if the Template has
     filter_by_deviceprofile=true it only reaches APs whose Device Profile is in the
     Template's deviceprofile_ids list, otherwise it reaches every AP at an applying site.
  2. Whether that AP has actually had a 5GHz client in the last N hours (one or more
     thresholds), via a targeted per-AP session lookup.

Cross-referencing the two flags up APs that publish a 5GHz SSID but show no real 5GHz
client activity — the interesting case that neither report alone surfaces.

Config file: mist_wlan_activity_report.ini
  [mist]       api_base, org_id, api_token
  [output]     directory
  [thresholds] hours (comma-separated, e.g. 24,48) — default set of windows to check;
               can still be overridden per-run with --hours on the command line.

Usage:
  python3 mist_wlan_activity_report.py
  python3 mist_wlan_activity_report.py --hours 12 24 48 72

Requirements:
  pip install requests openpyxl
"""
import argparse
import configparser
import os
import sys
import time
from collections import defaultdict
from datetime import datetime, timezone

import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

CONFIG_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "mist_wlan_activity_report.ini")
if not os.path.exists(CONFIG_PATH):
    sys.exit(
        f"Config file not found: {CONFIG_PATH}\n"
        f"Copy mist_wlan_activity_report.ini.example to mist_wlan_activity_report.ini and fill in your org details."
    )

config = configparser.ConfigParser()
config.read(CONFIG_PATH)

API_BASE = config.get("mist", "api_base")
ORG_ID = config.get("mist", "org_id")
API_TOKEN = config.get("mist", "api_token")
OUTPUT_DIR = os.path.expanduser(config.get("output", "directory"))

DEFAULT_HOURS = [
    int(h.strip()) for h in config.get("thresholds", "hours", fallback="24,48").split(",") if h.strip()
]

HEADERS = {"Authorization": f"Token {API_TOKEN}"}
RATE_HEADROOM = 10   # keep this many calls in reserve before pausing
HEADER_COLOR = "1565C0"


def format_eta(seconds):
    if seconds < 60:
        return f"{int(seconds)}s"
    m, s = divmod(int(seconds), 60)
    return f"{m}m {s}s"


# ---------------------------------------------------------------------------
# Rate-limit helpers (same pattern as mist_ap_details.py / mist_5ghz_client_check.py)
# ---------------------------------------------------------------------------

def fetch_api_usage():
    resp = requests.get(f"{API_BASE}/self/usage", headers=HEADERS)
    resp.raise_for_status()
    d = resp.json()
    return d.get("requests", 0), d.get("request_limit", 5000), d.get("seconds", 0)


def remaining_and_reset(used, limit):
    remaining = limit - used
    now = datetime.now()
    secs_into_hour = now.minute * 60 + now.second
    secs_left = max(0, 3600 - secs_into_hour)
    return remaining, secs_left


def wait_for_reset(secs_left, reason=""):
    wait = secs_left + 15
    if reason:
        print(f"\n  {reason}")
    print(f"  Waiting {wait // 60}m {wait % 60}s for the rate-limit window to reset...")
    deadline = time.time() + wait
    try:
        while True:
            left = int(deadline - time.time())
            if left <= 0:
                break
            m, s = divmod(left, 60)
            sys.stdout.write(f"\r  Resuming in {m}m {s:02d}s ...   ")
            sys.stdout.flush()
            time.sleep(1)
    except KeyboardInterrupt:
        print("\n\n  Wait interrupted by user. Exiting.")
        sys.exit(1)
    print("\r  Rate-limit window has reset. Continuing...          ")


def check_rate_limit(calls_needed, context=""):
    used, limit = fetch_api_usage()[:2]
    remaining, secs_left = remaining_and_reset(used, limit)

    if remaining - RATE_HEADROOM >= calls_needed:
        return remaining, secs_left

    m, s = divmod(secs_left, 60)
    label = f" ({context})" if context else ""
    print(f"\n\n  {'='*56}")
    print(f"  *** RATE LIMIT WARNING{label} ***")
    print(f"  Calls needed:    {calls_needed}")
    print(f"  Calls remaining: {remaining}  (limit {limit}/hour, used {used})")
    print(f"  Window resets in ~{m}m {s}s")
    print(f"  {'='*56}")
    print("  Options:")
    print("    w — wait for the window to reset then continue automatically")
    print("    q — quit now (no report will be saved)")

    while True:
        choice = input("  Your choice (w/q): ").strip().lower()
        if choice == "q":
            print("  Aborted by user.")
            sys.exit(0)
        if choice == "w":
            wait_for_reset(secs_left)
            used, limit = fetch_api_usage()[:2]
            remaining, secs_left = remaining_and_reset(used, limit)
            print(f"  New remaining calls: {remaining}")
            return remaining, secs_left
        print("  Please enter 'w' to wait or 'q' to quit.")


# ---------------------------------------------------------------------------
# API helpers
# ---------------------------------------------------------------------------

def fetch_org_info():
    resp = requests.get(f"{API_BASE}/orgs/{ORG_ID}", headers=HEADERS)
    resp.raise_for_status()
    return resp.json()


def fetch_org_stats():
    resp = requests.get(f"{API_BASE}/orgs/{ORG_ID}/stats", headers=HEADERS)
    resp.raise_for_status()
    return resp.json()


def fetch_all_org(path):
    """GET a plain-list org-level endpoint (templates/wlans/sitegroups are small at
    org scope, but page defensively anyway)."""
    items = []
    page = 1
    while True:
        url = f"{API_BASE}/orgs/{ORG_ID}/{path}?limit=1000&page={page}"
        resp = requests.get(url, headers=HEADERS)
        resp.raise_for_status()
        data = resp.json()
        if not data:
            break
        items.extend(data)
        if len(data) < 1000:
            break
        page += 1
    return items


def fetch_all_sites():
    """Return {site_id: site_name} for every site in the org, bulk-paginated."""
    site_map = {}
    page = 1
    api_calls = 0
    while True:
        url = f"{API_BASE}/orgs/{ORG_ID}/sites?limit=1000&page={page}"
        resp = requests.get(url, headers=HEADERS)
        resp.raise_for_status()
        api_calls += 1
        data = resp.json()
        if not data:
            break
        for site in data:
            site_map[site["id"]] = site.get("name", site["id"])
        sys.stdout.write(f"\r  Fetched {len(site_map)} sites (page {page})   ")
        sys.stdout.flush()
        if len(data) < 1000:
            break
        page += 1
    print(f"\r  Site fetch complete - {len(site_map)} sites resolved                ")
    return site_map, api_calls


def fetch_site_ap_stats(site_id):
    """Fetch AP stats for one site (assumes <=1000 APs per site). Site-level (rather
    than org-level) is required for uptime, port_stat, and radio_stat."""
    url = f"{API_BASE}/sites/{site_id}/stats/devices?type=ap&limit=1000"
    resp = requests.get(url, headers=HEADERS)
    resp.raise_for_status()
    return resp.json()


def fetch_aps_by_site(site_map):
    all_aps = []
    total = len(site_map)
    api_calls = 0
    call_times = []

    for idx, (site_id, site_name) in enumerate(site_map.items(), 1):
        if idx == 1 or idx % 50 == 0:
            calls_still_needed = total - idx + 1
            check_rate_limit(calls_still_needed, context=f"site {idx}/{total}")

        t0 = time.time()
        aps = fetch_site_ap_stats(site_id)
        call_times.append(time.time() - t0)
        api_calls += 1
        for ap in aps:
            ap["_site_name"] = site_name
        all_aps.extend(aps)

        avg_time = sum(call_times) / len(call_times)
        remaining_secs = (total - idx) * avg_time
        pct = idx / total * 100
        sys.stdout.write(f"\r  Site {idx}/{total} ({pct:.0f}%) | {len(all_aps)} APs so far | {avg_time:.2f}s/site | ETA: {format_eta(remaining_secs)}   ")
        sys.stdout.flush()

    print(f"\r  AP fetch complete - {len(all_aps)} APs across {total} sites                    ")
    return all_aps, api_calls


def fetch_last_5ghz_session(mac, start, end):
    """Return the timestamp (epoch seconds) of the most recent 5GHz session for one AP
    in [start, end], or None if it had no 5GHz client in that window.

    Deliberately targeted at a single AP (ap=<mac>, limit=1, sort=-timestamp) rather than
    paging through org-wide /clients/sessions/search: on a busy org that endpoint can return
    tens of millions of rows for a 48h window (roaming/reassociation churn), which makes
    full pagination impractical. One narrow query per AP bounds the work to O(APs)."""
    url = (
        f"{API_BASE}/orgs/{ORG_ID}/clients/sessions/search"
        f"?ap={mac}&band=5&start={start}&end={end}&limit=1&sort=-timestamp"
    )
    resp = requests.get(url, headers=HEADERS)
    resp.raise_for_status()
    results = resp.json().get("results", [])
    if not results:
        return None
    s = results[0]
    return s.get("disconnect") or s.get("timestamp") or s.get("connect")


def fetch_last_5ghz_by_ap(all_aps, start, end):
    last_seen = {}
    total = len(all_aps)
    api_calls = 0
    call_times = []

    for idx, ap in enumerate(all_aps, 1):
        mac = (ap.get("mac") or "").lower()
        if not mac:
            continue

        if idx == 1 or idx % 50 == 0:
            calls_still_needed = total - idx + 1
            check_rate_limit(calls_still_needed, context=f"AP {idx}/{total}")

        t0 = time.time()
        ts = fetch_last_5ghz_session(mac, start, end)
        call_times.append(time.time() - t0)
        api_calls += 1
        if ts is not None:
            last_seen[mac] = ts

        avg_time = sum(call_times) / len(call_times)
        remaining_secs = (total - idx) * avg_time
        pct = idx / total * 100
        sys.stdout.write(f"\r  AP {idx}/{total} ({pct:.0f}%) | {avg_time:.2f}s/call | ETA: {format_eta(remaining_secs)}   ")
        sys.stdout.flush()

    print(f"\r  5GHz lookup complete - {len(last_seen)}/{total} APs had a 5GHz client in the window                    ")
    return last_seen, api_calls


# ---------------------------------------------------------------------------
# WLAN Template resolution (same logic as mist_deviceprofile_wlan_report.py,
# applied per-AP instead of per (site, device profile) group)
# ---------------------------------------------------------------------------

def expand_sites(site_ids, sitegroup_ids, sitegroup_map):
    result = set(site_ids or [])
    for sg_id in (sitegroup_ids or []):
        result |= sitegroup_map.get(sg_id, set())
    return result


def build_effective_sites(templates, sitegroup_map):
    """Attach '_effective_sites' (applies minus exceptions) to each template in place."""
    for t in templates:
        applies = t.get("applies") or {}
        exceptions = t.get("exceptions") or {}
        applies_sites = expand_sites(applies.get("site_ids"), applies.get("sitegroup_ids"), sitegroup_map)
        exception_sites = expand_sites(exceptions.get("site_ids"), exceptions.get("sitegroup_ids"), sitegroup_map)
        t["_effective_sites"] = applies_sites - exception_sites


def template_reaches(template, site_id, deviceprofile_id):
    if site_id not in template["_effective_sites"]:
        return False
    if template.get("filter_by_deviceprofile"):
        return deviceprofile_id in (template.get("deviceprofile_ids") or [])
    return True


def format_rateset(band_key, wlan):
    bands = wlan.get("bands") or []
    if band_key not in bands:
        return ""
    band_cfg = (wlan.get("rateset") or {}).get(band_key) or {}
    template = band_cfg.get("template", "")
    legacy = band_cfg.get("legacy") or []
    min_rssi = band_cfg.get("min_rssi", 0)
    parts = [template] if template else []
    if legacy:
        parts.append(f"[{','.join(legacy)}]")
    if min_rssi:
        parts.append(f"min_rssi={min_rssi}")
    return " ".join(parts)


def resolve_wlans_for_ap(ap, templates, wlans_by_template):
    """Return a list of {ssid, template, bands, rates24, rates5, rates6} dicts for
    every WLAN that reaches this AP, based on its site + device profile."""
    site_id = ap.get("site_id") or ""
    dp_id = ap.get("deviceprofile_id") or ""
    reaching = []
    for t in templates:
        if not template_reaches(t, site_id, dp_id):
            continue
        for w in wlans_by_template.get(t["id"], []):
            reaching.append({
                "ssid": w.get("ssid", ""),
                "template": t.get("name", ""),
                "bands": w.get("bands") or [],
                "rates24": format_rateset("24", w),
                "rates5": format_rateset("5", w),
                "rates6": format_rateset("6", w),
            })
    return reaching


# ---------------------------------------------------------------------------
# Excel helpers
# ---------------------------------------------------------------------------

def style_header(ws, headers, fill_color=HEADER_COLOR):
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border


def auto_width(ws, headers):
    for col_idx, _ in enumerate(headers, 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for cell in ws[col_letter]:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 3, 60)


def finish_sheet(ws, headers):
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    ws.freeze_panes = "A2"
    auto_width(ws, headers)


def write_rows(ws, headers, rows):
    for r, row in enumerate(rows, 2):
        for c, h in enumerate(headers, 1):
            ws.cell(row=r, column=c, value=row.get(h, ""))


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="Report which WLANs reach each AP (via WLAN Template resolution) alongside "
                     "5GHz client activity for one or more hour thresholds."
    )
    parser.add_argument("--hours", "-x", type=int, nargs="+", default=None,
                         help=f"Hour thresholds to check, e.g. --hours 24 48 "
                              f"(default: from [thresholds] hours in the ini file, currently {DEFAULT_HOURS})")
    args = parser.parse_args()
    thresholds = sorted(set(args.hours)) if args.hours else sorted(set(DEFAULT_HOURS))
    max_hours = max(thresholds)

    now = time.time()
    window_start = int(now - max_hours * 3600)
    window_end = int(now)

    print("Fetching organisation info...")
    org_info = fetch_org_info()
    org_name = org_info.get("name", "Unknown")
    org_stats = fetch_org_stats()
    total_aps = org_stats.get("num_devices_connected", 0) + org_stats.get("num_devices_disconnected", 0)
    num_sites = org_stats.get("num_sites", 0)

    usage_used, usage_limit, _ = fetch_api_usage()
    rate_remaining, secs_left = remaining_and_reset(usage_used, usage_limit)
    m, s = divmod(secs_left, 60)

    est_site_list_pages = max(1, (num_sites // 1000) + 1)
    est_wlan_config_calls = 3   # templates + wlans + sitegroups (each typically 1 page)
    est_total_calls = est_site_list_pages + est_wlan_config_calls + num_sites + total_aps

    print(f"\n{'='*60}")
    print(f"  Mist WLAN Coverage + 5GHz Client Activity Report")
    print(f"{'='*60}")
    print(f"  Organisation:     {org_name}")
    print(f"  Org ID:           {ORG_ID}")
    print(f"  Total APs:        ~{total_aps}")
    print(f"  Total Sites:      ~{num_sites}")
    print(f"  Thresholds:       {', '.join(str(h) + 'h' for h in thresholds)}")
    print(f"{'='*60}")
    print(f"  This report will:")
    print(f"    1. Fetch WLAN templates, WLANs, site groups (~{est_wlan_config_calls} API calls)")
    print(f"    2. Fetch site list                          (~{est_site_list_pages} API calls)")
    print(f"    3. Fetch AP stats per site                   (~{num_sites} API calls, one per site)")
    print(f"    4. Look up last 5GHz session per AP          (~{total_aps} API calls, one per AP)")
    print(f"  Estimated total API calls: ~{est_total_calls}")
    print(f"{'='*60}")
    print(f"  API rate limit:    {usage_limit} calls/hour")
    print(f"  Used this hour:    {usage_used}")
    print(f"  Remaining:         {rate_remaining}  (window resets in ~{m}m {s}s)")

    if rate_remaining - RATE_HEADROOM < est_total_calls:
        windows_needed = -(-est_total_calls // max(1, usage_limit))  # ceil
        print(f"\n  *** This run needs more calls than one rate-limit window provides. ***")
        print(f"  Expect ~{windows_needed} wait-for-reset pause(s) (~1h each) before it completes.")
    print(f"{'='*60}")

    confirm = input("\n  Proceed? (y/n): ").strip().lower()
    if confirm != "y":
        print("  Aborted.")
        return

    start_time = time.time()

    print("\nFetching WLAN templates, WLANs and site groups...")
    templates = fetch_all_org("templates")
    wlans = fetch_all_org("wlans")
    sitegroups = fetch_all_org("sitegroups")
    print(f"  {len(templates)} templates, {len(wlans)} WLANs, {len(sitegroups)} site groups.")

    sitegroup_map = {sg["id"]: set(sg.get("site_ids") or []) for sg in sitegroups}
    build_effective_sites(templates, sitegroup_map)

    wlans_by_template = defaultdict(list)
    for w in wlans:
        tid = w.get("template_id")
        if tid:
            wlans_by_template[tid].append(w)

    print("\nFetching site list...")
    site_map, site_api_calls = fetch_all_sites()

    print(f"\nFetching AP stats per site ({len(site_map)} sites)...")
    all_aps, ap_api_calls = fetch_aps_by_site(site_map)
    print(f"Total APs fetched: {len(all_aps)}")

    print(f"\nLooking up last 5GHz session per AP (window: last {max_hours}h, covers all shorter thresholds too)...")
    last_seen, session_api_calls = fetch_last_5ghz_by_ap(all_aps, window_start, window_end)

    print("\nResolving WLAN coverage per AP...")
    ap_wlans = {}   # mac -> list of reaching WLAN dicts
    for ap in all_aps:
        mac = (ap.get("mac") or "").lower()
        ap_wlans[mac] = resolve_wlans_for_ap(ap, templates, wlans_by_template)

    ap_headers = [
        "Name", "Site Name", "Site ID", "Device Profile", "MAC", "Serial", "Model", "Firmware", "Status",
        "Uptime (hours)", "ETH0 RX Bytes", "ETH0 TX Bytes",
        "Band 5 Num WLANs", "Band 5 Channel", "Band 5 Bandwidth", "Band 5 Power",
        "Band 5 RX Bytes", "Band 5 TX Bytes",
        "Published SSIDs", "Published SSID Count",
        "Published 5GHz SSIDs", "Published 5GHz SSID Count",
    ]
    threshold_headers = [f"5GHz Clients ({h}h)" for h in thresholds]
    all_headers = ap_headers + ["Last 5GHz Client (UTC)", "Hours Since Last 5GHz Client"] + threshold_headers

    all_rows = []
    no_clients_rows = {h: [] for h in thresholds}
    detail_rows = []

    for ap in all_aps:
        mac = (ap.get("mac") or "").lower()
        ap_last_seen = last_seen.get(mac)
        hours_since = (now - ap_last_seen) / 3600 if ap_last_seen else None

        uptime = ap.get("uptime")
        eth0 = ap.get("port_stat", {}).get("eth0", {})
        band_5 = ap.get("radio_stat", {}).get("band_5", {})

        reaching = ap_wlans.get(mac, [])
        ssids = sorted({w["ssid"] for w in reaching if w["ssid"]})
        five_ghz = sorted({w["ssid"] for w in reaching if "5" in w["bands"] and w["ssid"]})

        row = {
            "Name": ap.get("name", ""),
            "Site Name": ap.get("_site_name", ""),
            "Site ID": ap.get("site_id", ""),
            "Device Profile": ap.get("deviceprofile_name") or ap.get("deviceprofile_id") or "",
            "MAC": mac,
            "Serial": ap.get("serial", ""),
            "Model": ap.get("model", ""),
            "Firmware": ap.get("version", ""),
            "Status": ap.get("status", ""),
            "Uptime (hours)": round(uptime / 3600, 1) if uptime is not None else "",
            "ETH0 RX Bytes": eth0.get("rx_bytes", ""),
            "ETH0 TX Bytes": eth0.get("tx_bytes", ""),
            "Band 5 Num WLANs": band_5.get("num_wlans", ""),
            "Band 5 Channel": band_5.get("channel", ""),
            "Band 5 Bandwidth": band_5.get("bandwidth", ""),
            "Band 5 Power": band_5.get("power", ""),
            "Band 5 RX Bytes": band_5.get("rx_bytes", ""),
            "Band 5 TX Bytes": band_5.get("tx_bytes", ""),
            "Published SSIDs": ", ".join(ssids),
            "Published SSID Count": len(ssids),
            "Published 5GHz SSIDs": ", ".join(five_ghz),
            "Published 5GHz SSID Count": len(five_ghz),
            "Last 5GHz Client (UTC)": (
                datetime.fromtimestamp(ap_last_seen, tz=timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
                if ap_last_seen else f"None in last {max_hours}h"
            ),
            "Hours Since Last 5GHz Client": round(hours_since, 1) if hours_since is not None else "",
        }

        for h in thresholds:
            has_client = hours_since is not None and hours_since <= h
            row[f"5GHz Clients ({h}h)"] = "Yes" if has_client else "No"
            if not has_client:
                no_clients_rows[h].append(row)

        all_rows.append(row)

        for w in reaching:
            detail_rows.append({
                "AP Name": ap.get("name", ""),
                "Site Name": ap.get("_site_name", ""),
                "Device Profile": row["Device Profile"],
                "MAC": mac,
                "SSID": w["ssid"],
                "WLAN Template": w["template"],
                "Bands": ", ".join(w["bands"]),
                "2.4GHz Rates": w["rates24"],
                "5GHz Rates": w["rates5"],
                "6GHz Rates": w["rates6"],
                "Last 5GHz Client (UTC)": row["Last 5GHz Client (UTC)"],
                **{f"5GHz Clients ({h}h)": row[f"5GHz Clients ({h}h)"] for h in thresholds},
            })

    quiet_all = [r for r in all_rows if all(r[f"5GHz Clients ({h}h)"] == "No" for h in thresholds)]
    # A disconnected AP trivially has no clients — that's not an RF/coverage issue, it's an
    # availability issue that belongs in a different report. Only a *connected* AP that
    # publishes a 5GHz SSID and still shows no 5GHz client is the interesting mismatch.
    quiet_connected = [r for r in quiet_all if r["Status"] == "connected"]
    possible_issues = [r for r in quiet_connected if r["Published 5GHz SSID Count"] > 0]

    # No WLAN Template reaches this AP at all (any band) — a config/coverage gap, independent
    # of whether it's ever had a client. Worth its own sheet rather than folding into
    # Possible Issues (which is specifically about the 5GHz-published-but-quiet mismatch).
    no_wlans_published = [r for r in all_rows if r["Published SSID Count"] == 0]
    no_wlans_connected = [r for r in no_wlans_published if r["Status"] == "connected"]

    print(f"\nCategorised: {len(all_rows)} APs total")
    for h in thresholds:
        print(f"  No 5GHz clients in last {h}h:  {len(no_clients_rows[h])}")
    print(f"  No 5GHz clients in ANY of {', '.join(str(h)+'h' for h in thresholds)}: {len(quiet_all)}")
    print(f"    - of those, disconnected (expected, not an issue): {len(quiet_all) - len(quiet_connected)}")
    print(f"    - of those, connected but still quiet:             {len(quiet_connected)}")
    print(f"  Possible issues (connected, publishes >=1 5GHz SSID, no client): {len(possible_issues)}")
    print(f"  No WLANs published at all (any band):            {len(no_wlans_published)} "
          f"({len(no_wlans_connected)} connected, {len(no_wlans_published) - len(no_wlans_connected)} disconnected)")

    print("\nBuilding Excel spreadsheet...")
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "All APs"
    style_header(ws1, all_headers)
    write_rows(ws1, all_headers, all_rows)
    finish_sheet(ws1, all_headers)

    no_client_headers = ap_headers + ["Last 5GHz Client (UTC)", "Hours Since Last 5GHz Client"]
    colors = ["E65100", "B71C1C", "6A1B9A", "37474F"]
    for i, h in enumerate(thresholds):
        ws = wb.create_sheet(f"No 5GHz Clients - {h}h")
        style_header(ws, no_client_headers, colors[i % len(colors)])
        write_rows(ws, no_client_headers, no_clients_rows[h])
        finish_sheet(ws, no_client_headers)

    ws_quiet = wb.create_sheet("Quiet in ALL Thresholds")
    style_header(ws_quiet, no_client_headers, "212121")
    write_rows(ws_quiet, no_client_headers, quiet_all)
    finish_sheet(ws_quiet, no_client_headers)

    ws_issues = wb.create_sheet("Possible Issues")
    style_header(ws_issues, no_client_headers, "C62828")
    write_rows(ws_issues, no_client_headers, possible_issues)
    finish_sheet(ws_issues, no_client_headers)

    ws_no_wlans = wb.create_sheet("No WLANs Published")
    style_header(ws_no_wlans, no_client_headers, "4E342E")
    write_rows(ws_no_wlans, no_client_headers, no_wlans_published)
    finish_sheet(ws_no_wlans, no_client_headers)

    detail_headers = [
        "AP Name", "Site Name", "Device Profile", "MAC", "SSID", "WLAN Template", "Bands",
        "2.4GHz Rates", "5GHz Rates", "6GHz Rates", "Last 5GHz Client (UTC)",
    ] + threshold_headers
    ws_detail = wb.create_sheet("AP x SSID Detail")
    style_header(ws_detail, detail_headers, "1A237E")
    write_rows(ws_detail, detail_headers, detail_rows)
    finish_sheet(ws_detail, detail_headers)

    timestamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    safe_org_name = "".join(c if c.isalnum() or c in (" ", "-", "_") else "_" for c in org_name).strip().replace(" ", "_")
    filename = f"Mist_WLAN_Activity_Report_{safe_org_name}_{timestamp}.xlsx"
    filepath = os.path.join(OUTPUT_DIR, filename)
    wb.save(filepath)

    total_api = site_api_calls + ap_api_calls + session_api_calls + est_wlan_config_calls
    elapsed = time.time() - start_time
    print(f"\n{'='*60}")
    print(f"  Mist WLAN Activity Report Summary - {org_name}")
    print(f"{'='*60}")
    print(f"  Total APs processed:            {len(all_aps)}")
    for h in thresholds:
        print(f"  No 5GHz clients ({h}h):           {len(no_clients_rows[h])}")
    print(f"  Quiet in ALL thresholds:         {len(quiet_all)} ({len(quiet_all) - len(quiet_connected)} disconnected, {len(quiet_connected)} connected)")
    print(f"  Possible issues (connected, SSID, no client): {len(possible_issues)}")
    print(f"  No WLANs published (any band):   {len(no_wlans_published)} ({len(no_wlans_connected)} connected)")
    print(f"  AP x SSID detail rows:           {len(detail_rows)}")
    print(f"{'='*60}")
    print(f"  API calls - WLAN config:         {est_wlan_config_calls}")
    print(f"  API calls - Site list:           {site_api_calls}")
    print(f"  API calls - AP stats/site:       {ap_api_calls}")
    print(f"  API calls - Per-AP 5GHz lookups: {session_api_calls}")
    print(f"  Total API calls:                 {total_api}")
    print(f"  Total elapsed time:              {format_eta(elapsed)}")
    print(f"{'='*60}")
    print(f"  Report saved: {filepath}")


if __name__ == "__main__":
    main()
